"""
backend/parsers/spreadsheet_parser.py
=====================================
Parser for spreadsheet workbooks: XLSX, XLSM, XLS, ODS.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from backend.parsers.base import BaseParser, CorruptedFileError, EncryptedFileError, ParsedDocument, ParsedField


class SpreadsheetParser(BaseParser):
    """Parses Excel (XLSX, XLS) and OpenDocument (ODS) spreadsheets."""

    SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".xls", ".ods"}

    def __init__(self, max_sample_rows: int = 1000) -> None:
        self.max_sample_rows = max_sample_rows

    def can_parse(self, ext: str, mime: Optional[str] = None, header_bytes: Optional[bytes] = None) -> bool:
        return ext.lower() in self.SUPPORTED_EXTENSIONS

    def parse(self, path: Path) -> ParsedDocument:
        path = Path(path)
        file_name = path.name
        ext = path.suffix.lower()
        file_type = "xlsx" if ext in {".xlsx", ".xlsm"} else ext.lstrip(".")

        if ext in {".xlsx", ".xlsm"}:
            return self._parse_openpyxl(path, file_name, file_type)
        elif ext == ".ods":
            return self._parse_ods(path, file_name, file_type)
        else:
            return self._parse_fallback(path, file_name, file_type)

    def _parse_openpyxl(self, path: Path, file_name: str, file_type: str) -> ParsedDocument:
        try:
            import openpyxl
        except ImportError:
            raise CorruptedFileError("openpyxl library is required to parse Excel files.", file_name)

        try:
            # data_only=True evaluates formulas to cached values
            wb = openpyxl.load_workbook(filename=str(path), read_only=True, data_only=True)
        except Exception as e:
            err_msg = str(e).lower()
            if "encrypted" in err_msg or "password" in err_msg:
                raise EncryptedFileError(f"Spreadsheet '{file_name}' is password-protected or encrypted.", file_name)
            raise CorruptedFileError(f"Corrupted or invalid Excel file '{file_name}': {str(e)}", file_name)

        all_fields: list[ParsedField] = []
        total_rows = 0
        sheet_names = wb.sheetnames

        try:
            for sheet_name in sheet_names:
                ws = wb[sheet_name]
                rows_iter = ws.iter_rows(values_only=True)
                
                # Find header row
                header = None
                for candidate in rows_iter:
                    if candidate and any(c is not None and str(c).strip() != "" for c in candidate):
                        header = candidate
                        break
                
                if not header:
                    continue

                clean_headers = []
                for i, h in enumerate(header):
                    if h is not None and str(h).strip():
                        name = re.sub(r"\s+", "_", str(h).strip())
                    else:
                        name = f"col_{i+1}"
                    clean_headers.append(name)

                # Prefix sheet name only if workbook has multiple sheets
                prefix = f"{sheet_name}:" if len(sheet_names) > 1 else ""

                col_data: dict[str, list[str]] = {h: [] for h in clean_headers}
                s_rows = 0

                for row in rows_iter:
                    if not row or not any(row):
                        continue
                    s_rows += 1
                    for idx, val in enumerate(row):
                        if idx < len(clean_headers) and val is not None:
                            s_val = str(val).strip()
                            if s_val:
                                col_data[clean_headers[idx]].append(s_val)
                    if s_rows >= self.max_sample_rows:
                        break

                total_rows += s_rows
                for h in clean_headers:
                    display_name = f"{prefix}{h}"
                    all_fields.append(
                        ParsedField(
                            field_name=display_name,
                            values=col_data[h],
                            field_path=f"{sheet_name}/{h}",
                        )
                    )
        finally:
            wb.close()

        return ParsedDocument(
            file_name=file_name,
            file_type=file_type,
            fields=all_fields,
            total_records=total_rows,
            metadata={"sheets": sheet_names, "total_fields": len(all_fields)},
        )

    def _parse_ods(self, path: Path, file_name: str, file_type: str) -> ParsedDocument:
        try:
            import pandas as pd
            excel_file = pd.ExcelFile(path, engine="odf")
            all_fields: list[ParsedField] = []
            total_rows = 0
            for sheet_name in excel_file.sheet_names:
                df = excel_file.parse(sheet_name, nrows=self.max_sample_rows)
                prefix = f"{sheet_name}:" if len(excel_file.sheet_names) > 1 else ""
                total_rows += len(df)
                for col in df.columns:
                    col_str = re.sub(r"\s+", "_", str(col).strip())
                    vals = [str(v).strip() for v in df[col].dropna().tolist() if str(v).strip()]
                    all_fields.append(
                        ParsedField(
                            field_name=f"{prefix}{col_str}",
                            values=vals,
                            field_path=f"{sheet_name}/{col_str}",
                        )
                    )
            return ParsedDocument(
                file_name=file_name,
                file_type="ods",
                fields=all_fields,
                total_records=total_rows,
                metadata={"sheets": excel_file.sheet_names},
            )
        except Exception as e:
            raise CorruptedFileError(f"Failed to parse ODS file '{file_name}': {str(e)}", file_name)

    def _parse_fallback(self, path: Path, file_name: str, file_type: str) -> ParsedDocument:
        try:
            import pandas as pd
            df = pd.read_excel(path, nrows=self.max_sample_rows)
            fields = []
            for col in df.columns:
                col_name = re.sub(r"\s+", "_", str(col).strip())
                vals = [str(v).strip() for v in df[col].dropna().tolist() if str(v).strip()]
                fields.append(ParsedField(field_name=col_name, values=vals, field_path=col_name))
            return ParsedDocument(
                file_name=file_name,
                file_type=file_type,
                fields=fields,
                total_records=len(df),
            )
        except Exception as e:
            raise CorruptedFileError(f"Failed to read spreadsheet '{file_name}': {str(e)}", file_name)
